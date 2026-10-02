import json
from observation_fixtures import fixture_observations, session_runtime, run_turn, extraction_steps
from datetime import date
from decimal import Decimal
from io import BytesIO
from unittest.mock import patch

import httpx
import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook

from valuationagent.api.main import create_app
from valuationagent.application.research import ResearchService
from valuationagent.application.research_valuation import (
    ResearchValuationAssembler,
    _period,
)
from valuationagent.core.documents import parse_document
from valuationagent.core.tools import NoArguments, ToolSpec
from valuationagent.llm.client import LlmError, OpenAICompatibleClient
from valuationagent.schemas.agent import SearchQuery, SearchResult
from valuationagent.schemas.models import ModelConnectionInput
from valuationagent.schemas.research import DocumentSummary, FactCandidate, ResearchTurn
from valuationagent.search.providers import MockSearchProvider
from valuationagent.storage.sqlite import SQLiteRunStore


class ScriptedModel:
    def __init__(self, actions):
        self.actions = iter([*actions, ("finish_response", {"answer": "本轮测试工具执行完毕。"})])
        self.calls = []
        self.kwargs = []

    def chat(self, messages, **kwargs):
        self.calls.append(json.loads(json.dumps(messages)))
        self.kwargs.append(kwargs)
        name, args = next(self.actions)
        return {"tool_calls": [{"id": f"call_{len(self.calls)}", "type": "function",
            "function": {"name": name, "arguments": json.dumps(args, ensure_ascii=False)}}],
            "reasoning_content": "PRIVATE_TEST_REASONING"}


@pytest.fixture
def service(tmp_path):
    return ResearchService(SQLiteRunStore(tmp_path))


def candidate(file_id, **changes):
    return {"metric": "revenue", "raw_value": "12000", "unit": "万元", "period": "2025",
        "scope": "consolidated", "role": "historical", "block_id": file_id + ":1",
        "quote": "2025 年合并报表，单位万元。营业收入：12000。", **changes}


def upload(store):
    return store.save_upload("年报摘录.txt", "historical_financials", "text/plain",
        "2025 年合并报表，单位万元。营业收入：12000。".encode())["file_id"]


def ready_online_session(service, *, llm=None):
    session = service.create(llm=llm)
    session.draft.company = "测试股份"
    session.draft.ticker = "600000.SH"
    session.draft.valuation_date = date(2026, 9, 24)
    session.draft.methods = ["dcf", "pe", "ev_ebitda"]
    session.data_source_preference = "online"
    service.store.save_research(session)
    return session


def test_unavailable_search_provider_closes_without_three_pointless_retries(service):
    model = ScriptedModel([
        ("search_sources", {
            "query": "600000 2024 财务字段",
            "reason": "补齐DCF基期",
            "purpose": "financials",
        }),
        ("finish_response", {
            "answer": "当前进程没有可用的网页搜索服务。",
            "outcome": "insufficient_data",
        }),
    ])
    session = service.create(llm=model, data_source_preference="web")
    session.pending_action = "valuation"
    session.draft = session.draft.model_copy(update={
        "company": "测试股份",
        "ticker": "600000.SH",
        "industry": "电子",
        "valuation_date": date(2026, 9, 24),
        "methods": ["dcf"],
    })
    service.store.save_research(session)

    state = service.turn(
        session.session_id,
        ResearchTurn(content="使用公开资料完成估值"),
    )

    assert len(model.calls) == 2
    assert len(state["session"]["search_history"]) == 1
    assert state["session"]["search_history"][0]["status"] == "not_configured"
    assert state["session"]["search_history"][0]["target_ticker"] == "600000.SH"
    assert state["session"]["outcome_status"] == "insufficient_data"
    assert "question" not in state["session"]


def _confirmed_statement_fact(metric, value, *, period="2024", unit="元", suffix=""):
    return FactCandidate(
        fact_id=f"fact_{metric}_{period}_{suffix or 'base'}",
        metric=metric,
        raw_value=str(value),
        unit=unit,
        normalized_value=str(value),
        period=period,
        scope="consolidated",
        block_id=f"message:{metric}_{period}_{suffix or 'base'}",
        quote=f"{metric} {value}",
        status="confirmed",
    )


def test_financial_period_parser_distinguishes_annual_ranges_and_interims():
    assert _period("2024年1-12月") == date(2024, 12, 31)
    assert _period("2024-01-01 至 2024-12-31") == date(2024, 12, 31)
    assert _period("2024年度") == date(2024, 12, 31)
    assert _period("2024年第三季度") is None
    assert _period("2024H1") is None


def test_structured_assembler_deterministically_derives_annual_report_metrics(service):
    session = service.create()
    raw_values = {
        "revenue": "1000",
        "profit_before_tax": "100",
        "income_tax_expense": "15",
        "interest_expense": "10",
        "depreciation_fixed_assets": "20",
        "amortization_intangible_assets": "5",
        "amortization_long_term_deferred_expenses": "2",
        # Synthetic source explicitly discloses zero ROU amortization; a
        # missing value must not be interpreted as zero for this lease balance.
        "depreciation_right_of_use": "0",
        "cash_paid_for_ppe_intangibles": "40",
        "inventory_decrease": "-2",
        "operating_receivables_decrease": "-3",
        "operating_payables_increase": "4",
        "cash_and_non_operating_assets": "200",
        "short_term_borrowings": "50",
        "current_portion_non_current_liabilities": "10",
        "long_term_borrowings": "60",
        "bonds_payable": "20",
        "lease_liabilities": "5",
        "common_shares": "100",
        "net_income_parent": "80",
    }
    session.facts = [
        _confirmed_statement_fact(
            metric,
            value,
            unit="股" if metric == "common_shares" else "元",
        )
        for metric, value in raw_values.items()
    ]

    snapshot = ResearchValuationAssembler()._structured_financials(session)[0]

    assert snapshot.ebit_margin == Decimal("0.11")
    assert snapshot.tax_rate == Decimal("0.15")
    assert snapshot.depreciation_amortization == Decimal(27)
    assert snapshot.capital_expenditure == Decimal(40)
    assert snapshot.change_operating_nwc == Decimal(1)
    assert snapshot.interest_bearing_debt == Decimal(145)
    assert snapshot.ebitda == Decimal(137)
    assert snapshot.calculation_methods["tax_rate"] == (
        "income_tax_expense / profit_before_tax"
    )
    assert snapshot.calculation_methods["change_operating_nwc"].startswith("-(")
    assert snapshot.statement_items["profit_before_tax"] == Decimal(100)
    assert {
        evidence.evidence_id for evidence in snapshot.evidence["ebitda"]
    } == {
        "fact_profit_before_tax_2024_base",
        "fact_interest_expense_2024_base",
        "fact_depreciation_fixed_assets_2024_base",
        "fact_amortization_intangible_assets_2024_base",
        "fact_amortization_long_term_deferred_expenses_2024_base",
        "fact_depreciation_right_of_use_2024_base",
    }
    assert "确定性推导" in snapshot.source_label


def test_structured_assembler_blocks_conflicting_confirmed_aliases(service):
    session = service.create()
    session.facts = [
        _confirmed_statement_fact("revenue", "1000", suffix="a"),
        _confirmed_statement_fact("营业收入", "2000", suffix="b"),
    ]

    error = ResearchValuationAssembler().structured_readiness_error(session)

    assert "同期间同口径冲突" in error
    assert "未静默覆盖" in error
    assert "1000 与 2000" in error


def test_structured_assembler_does_not_clamp_anomalous_effective_tax_rate(service):
    session = service.create()
    required_direct = {
        "revenue": "1000",
        "ebit_margin": "0.1",
        "depreciation_amortization": "20",
        "capital_expenditure": "30",
        "change_operating_nwc": "5",
        "cash_and_non_operating_assets": "200",
        "interest_bearing_debt": "100",
        "common_shares": "100",
        "net_income_parent": "80",
        "ebitda": "120",
        "profit_before_tax": "100",
        "income_tax_expense": "80",
    }
    session.facts = [
        _confirmed_statement_fact(
            metric,
            value,
            unit=("ratio" if metric == "ebit_margin" else "股" if metric == "common_shares" else "元"),
        )
        for metric, value in required_direct.items()
    ]

    error = ResearchValuationAssembler().structured_readiness_error(session)

    assert "所得税率" in error
    assert "0.6" not in error


def test_structured_assembler_records_passed_financial_reconciliations(service):
    session = service.create()
    values = {
        "revenue": "1000",
        "ebit": "100",
        "ebit_margin": "0.1",
        "profit_before_tax": "100",
        "income_tax_expense": "15",
        "tax_rate": "0.15",
        "depreciation_amortization": "20",
        "capital_expenditure": "30",
        "change_operating_nwc": "5",
        "cash_and_non_operating_assets": "200",
        "interest_bearing_debt": "100",
        "common_shares": "100",
        "net_income_parent": "80",
        "ebitda": "120",
    }
    session.facts = [
        _confirmed_statement_fact(
            metric,
            value,
            unit=(
                "ratio" if metric in {"ebit_margin", "tax_rate"}
                else "股" if metric == "common_shares"
                else "元"
            ),
        )
        for metric, value in values.items()
    ]

    snapshot = ResearchValuationAssembler()._structured_financials(session)[0]

    assert snapshot.calculation_methods["reconciliation.ebit_margin"].startswith("passed:")
    assert snapshot.calculation_methods["reconciliation.tax_rate"].startswith("passed:")
    assert snapshot.calculation_methods["reconciliation.ebitda"].startswith("passed:")


def test_structured_assembler_blocks_financial_reconciliation_conflicts(service):
    session = service.create()
    session.facts = [
        _confirmed_statement_fact("profit_before_tax", "100"),
        _confirmed_statement_fact("income_tax_expense", "30"),
        _confirmed_statement_fact("tax_rate", "0.15", unit="ratio"),
    ]

    error = ResearchValuationAssembler().structured_readiness_error(session)

    assert "财务勾稽冲突" in error
    assert "所得税率直接值为 0.15" in error
    assert "复算为 0.3" in error


def test_one_complete_year_is_blocked_before_formal_model_handoff(service):
    session = service.create()
    session.draft = session.draft.model_copy(update={
        "company": "测试汽车",
        "industry": "汽车制造",
        "valuation_date": date(2025, 4, 30),
        "methods": ["dcf"],
    })
    session.data_source_preference = "web"
    values = {
        "revenue": "1000",
        "ebit_margin": "0.1",
        "tax_rate": "0.15",
        "depreciation_amortization": "20",
        "capital_expenditure": "30",
        "change_operating_nwc": "5",
        "cash_and_non_operating_assets": "200",
        "interest_bearing_debt": "100",
        "common_shares": "100",
        "net_income_parent": "80",
        "ebitda": "120",
    }
    session.facts = [
        _confirmed_statement_fact(
            metric,
            value,
            unit=(
                "ratio" if metric in {"ebit_margin", "tax_rate"}
                else "股" if metric == "common_shares"
                else "元"
            ),
        )
        for metric, value in values.items()
    ]

    assembler = ResearchValuationAssembler()
    error = assembler.structured_readiness_error(session)

    assert "至少需要4个连续年度完整快照" in error
    assert "当前只有 1 个完整年度（2024）" in error
    assert "2021、2022、2023" in error
    with pytest.raises(ValueError, match="至少需要4个连续年度"):
        assembler.build(session)


def test_newer_partial_year_is_not_hidden_by_an_older_complete_snapshot(service):
    session = service.create()
    complete_values = {
        "revenue": "900",
        "ebit_margin": "0.1",
        "tax_rate": "0.15",
        "depreciation_amortization": "20",
        "capital_expenditure": "30",
        "change_operating_nwc": "5",
        "cash_and_non_operating_assets": "200",
        "interest_bearing_debt": "100",
        "common_shares": "100",
        "net_income_parent": "80",
        "ebitda": "120",
    }
    session.facts = [
        _confirmed_statement_fact(
            metric,
            value,
            period="2023",
            unit=(
                "ratio" if metric in {"ebit_margin", "tax_rate"}
                else "股" if metric == "common_shares"
                else "元"
            ),
        )
        for metric, value in complete_values.items()
    ]
    session.facts.append(
        _confirmed_statement_fact("revenue", "1000", period="2024")
    )

    error = ResearchValuationAssembler().structured_readiness_error(session)

    assert "优先补齐最近年度 2024-12-31" in error
    assert "归母净利润" in error


def test_search_snippet_financial_fact_requires_repair_not_blanket_confirmation(service):
    runtime = session_runtime(service)
    args = fixture_observations(runtime)
    args["file_id"] = "web_search_lead"
    for anchor in args["anchors"].values():
        anchor["block_id"] = "web_search_lead:1"
    runtime.session.documents.append(DocumentSummary(file_id="web_search_lead", name="搜索摘要", role="evidence",
        block_count=1, provenance_type="search_snippet", authority_tier="E"))
    service.store.save_research_blocks(runtime.session.session_id, "web_search_lead", [{
        "block_id": "web_search_lead:1", "file_id": "web_search_lead", "text": "检索摘要提及收入12000",
        "location": {"source_type": "web_search", "url": "https://example.test/report"}}])
    state, model = run_turn(runtime, [("extract_observations", args),
        ("finish_response", {"answer": "只有摘要尚未取得原文，不提取数值。", "outcome": "insufficient_data"})])
    assert not state["session"]["facts"]
    assert len(model.calls) == 2
    assert not service.store.list_runs()
    assert any(event["tool"] == "extract_observations" and event["payload"].get("output", {}).get("ok") is False
               for event in state["events"])

def test_repeated_partial_fact_batch_preserves_valid_candidate_without_recovery(service):
    runtime = session_runtime(service)
    args = fixture_observations(runtime, [{"metric": "net_income_parent", "raw_value": "1800", "unit": "万元"}])
    state, model = run_turn(runtime, [("extract_observations", args), ("extract_observations", args), *extraction_steps(args)])
    assert len(model.calls) == 6
    assert len(state["session"]["facts"]) == 1
    assert state["session"]["facts"][0]["status"] == "confirmed"
    assert state["session"]["facts"][0]["standard_metric"] == "net_income_parent"
    assert state["session"]["last_issue"] is None
    assert not service.store.list_runs()

def test_official_search_pdf_can_be_downloaded_parsed_and_traced(tmp_path):
    from reportlab.pdfgen import canvas

    output = BytesIO()
    pdf = canvas.Canvas(output)
    pdf.drawString(72, 720, "2023 consolidated revenue 12000 CNY")
    pdf.save()

    def handler(request):
        assert request.url.host == "static.cninfo.com.cn"
        return httpx.Response(
            200,
            headers={"content-type": "application/pdf"},
            content=output.getvalue(),
        )

    service = ResearchService(
        SQLiteRunStore(tmp_path / "remote-pdf"),
        remote_transport=httpx.MockTransport(handler),
    )
    session = service.create()
    lead_id = "web_cninfo_report"
    service.store.save_research_blocks(session.session_id, lead_id, [{
        "block_id": lead_id + ":1",
        "text": "annual report",
        "location": {
            "source_type": "web_search",
            "url": "https://static.cninfo.com.cn/finalpage/report.pdf",
            "provider": "mock-search",
            "search_query": "annual report",
        },
    }])
    session.documents.append(DocumentSummary(
        file_id=lead_id,
        name="Annual report lead",
        role="evidence",
        block_count=1,
        warnings=["联网搜索摘要；形成关键事实前应打开原始URL核对全文。"],
    ))

    result = service._fetch_search_source(session, lead_id)
    blocks = service.store.research_blocks(session.session_id, result["file_id"])

    assert result["status"] == "fetched"
    assert "revenue 12000" in blocks[0]["text"]
    assert blocks[0]["location"]["source_type"] == "remote_document"
    assert blocks[0]["location"]["parent_search_file_id"] == lead_id
    assert blocks[0]["location"]["source_url"].startswith("https://static.cninfo.com.cn/")
    document = next(item for item in session.documents if item.file_id == result["file_id"])
    assert document.provenance_type == "official_filing"
    assert document.authority_tier == "A"
    assert document.source_confidence == .99


def test_hkex_pdf_is_treated_as_official_exchange_disclosure(tmp_path):
    from reportlab.pdfgen import canvas

    output = BytesIO()
    pdf = canvas.Canvas(output)
    pdf.drawString(72, 720, "HKEX final allotment result")
    pdf.save()

    def handler(request):
        assert request.url.host == "www.hkexnews.hk"
        return httpx.Response(
            200,
            headers={"content-type": "application/pdf"},
            content=output.getvalue(),
        )

    service = ResearchService(
        SQLiteRunStore(tmp_path / "hkex-pdf"),
        remote_transport=httpx.MockTransport(handler),
    )
    session = service.create()
    lead_id = "web_hkex_allotment"
    service.store.save_research_blocks(session.session_id, lead_id, [{
        "block_id": lead_id + ":1",
        "text": "final allotment result",
        "location": {
            "source_type": "web_search",
            "url": "https://www.hkexnews.hk/listedco/final.pdf",
            "provider": "tavily",
            "published_at": "2025-06-18",
            "search_query": "issuer final allotment",
        },
    }])
    session.documents.append(DocumentSummary(
        file_id=lead_id,
        name="Final allotment lead",
        role="evidence",
        block_count=1,
    ))

    result = service._fetch_search_source(session, lead_id)
    blocks = service.store.research_blocks(session.session_id, result["file_id"])

    assert result["status"] == "fetched"
    assert blocks[0]["location"]["source_type"] == "remote_document"
    assert blocks[0]["location"]["source_domain"] == "www.hkexnews.hk"
    assert blocks[0]["location"]["published_at"] == "2025-06-18"
    assert not any("公开网页原文" in warning for warning in result["warnings"])
    document = next(item for item in session.documents if item.file_id == result["file_id"])
    assert document.provenance_type == "official_filing"
    assert document.authority_tier == "A"


def test_public_search_html_can_be_downloaded_read_and_traced(tmp_path):
    page = b"""<!doctype html><html><head><style>hidden</style></head><body>
    <h1>Industry policy</h1><p>Effective from 2025-01-01.</p>
    <table><tr><th>Measure</th><th>Value</th></tr><tr><td>Subsidy</td><td>10%</td></tr></table>
    <script>stealSecrets()</script></body></html>"""

    def handler(request):
        assert request.url == httpx.URL("https://example.com/policy")
        return httpx.Response(200, headers={"content-type": "text/html; charset=utf-8"}, content=page)

    service = ResearchService(
        SQLiteRunStore(tmp_path / "remote-html"),
        remote_transport=httpx.MockTransport(handler),
    )
    session = service.create()
    lead_id = "web_public_policy"
    service.store.save_research_blocks(session.session_id, lead_id, [{
        "block_id": lead_id + ":1",
        "text": "policy lead",
        "location": {
            "source_type": "web_search",
            "url": "https://example.com/policy",
            "provider": "tavily",
            "search_query": "industry policy",
        },
    }])
    session.documents.append(DocumentSummary(
        file_id=lead_id,
        name="Policy lead",
        role="evidence",
        block_count=1,
    ))

    result = service._fetch_search_source(session, lead_id)
    blocks = service.store.research_blocks(session.session_id, result["file_id"])
    extracted = "\n".join(block["text"] for block in blocks)

    assert result["status"] == "fetched"
    assert "Industry policy" in extracted
    assert "Subsidy" in extracted and "10%" in extracted
    assert "stealSecrets" not in extracted and "hidden" not in extracted
    assert blocks[0]["location"]["source_type"] == "remote_web_document"
    assert blocks[0]["location"]["parent_search_file_id"] == lead_id
    assert any("公开网页原文" in warning for warning in result["warnings"])
    document = next(item for item in session.documents if item.file_id == result["file_id"])
    assert document.provenance_type == "public_web"
    assert document.authority_tier == "C"
    assert document.source_confidence == .55


def test_public_search_fetch_rejects_private_network_targets(tmp_path):
    service = ResearchService(
        SQLiteRunStore(tmp_path / "blocked-web"),
        remote_transport=httpx.MockTransport(lambda request: httpx.Response(200)),
    )
    session = service.create()
    lead_id = "web_private_target"
    service.store.save_research_blocks(session.session_id, lead_id, [{
        "block_id": lead_id + ":1",
        "text": "unsafe lead",
        "location": {
            "source_type": "web_search",
            "url": "https://127.0.0.1/internal",
            "provider": "tavily",
        },
    }])
    session.documents.append(DocumentSummary(
        file_id=lead_id,
        name="Unsafe lead",
        role="evidence",
        block_count=1,
    ))

    with pytest.raises(ValueError, match="私有、回环或保留地址"):
        service._fetch_search_source(session, lead_id)


def test_web_financial_search_uses_official_catalogue_without_tavily_key(tmp_path):
    class OfficialCatalogue:
        def search_annual_reports(self, ticker, years, *, cutoff, company_name):
            assert ticker == "600000.SH"
            assert years == [2022, 2023, 2024]
            return SearchResult(
                query=SearchQuery(
                    query="600000 2022 2023 2024 年度报告",
                    ticker=ticker,
                    purpose="financials",
                ),
                provider="cninfo-announcements",
                provider_version="test",
                status="completed",
                hits=[{
                    "source_id": "cninfo_report2024",
                    "title": "测试股份2024年年度报告",
                    "url": "https://static.cninfo.com.cn/finalpage/2025-03-31/report.PDF",
                    "domain": "static.cninfo.com.cn",
                    "snippet": "巨潮资讯官方公告目录；证券代码 600000；报告年度 2024。",
                    "published_at": "2025-03-31",
                    "relevance": 1,
                }],
            )

    model = ScriptedModel([
        ("search_sources", {
            "query": "600000 2024/2023/2022 年度报告",
            "reason": "补齐历史年报",
            "purpose": "financials",
        }),
        ("finish_response", {
            "answer": "已从巨潮资讯官方公告目录定位年报，下一步下载原文。",
        }),
    ])
    service = ResearchService(
        SQLiteRunStore(tmp_path / "official-catalogue"),
        official_search_provider=OfficialCatalogue(),
    )
    session = ready_online_session(service, llm=model)
    session.data_source_preference = "web"
    service.store.save_research(session)

    state = service.turn(session.session_id, ResearchTurn(content="继续查找这三年的官方年报"))

    assert "question" not in state["session"]
    assert any(doc["file_id"] == "web_cninfo_report2024" for doc in state["session"]["documents"])
    lead = next(doc for doc in state["session"]["documents"] if doc["file_id"] == "web_cninfo_report2024")
    assert lead["provenance_type"] == "official_index"
    assert lead["authority_tier"] == "A"
    assert "官方公告目录" in state["messages"][-1]["content"]
    assert len(model.calls) == 2


def test_candidate_input_tolerates_common_scope_labels_and_adjacent_columns(service):
    runtime = session_runtime(service)
    args = fixture_observations(runtime, [{"metric": "capital_expenditure", "raw_value": "1,286,898,447.55"}])
    state, _ = run_turn(runtime, extraction_steps(args))
    fact = state["session"]["facts"][0]
    assert fact["scope"] == "consolidated" and fact["unit"] == "元"
    assert fact["normalized_value"] == "1286898447.55"
    assert not fact["warnings"]
    assert fact["verification"]["reading_proof"]["anchors"]["row0_value"]["quote"] == "1,286,898,447.55"

def test_failed_search_stops_before_a_second_llm_call(service):
    class FailedSearch:
        provider_id = "failed-test"
        version = "1"

        def search(self, query):
            return SearchResult(
                query=query,
                provider=self.provider_id,
                status="failed",
                error_code="SEARCH_AUTH_FAILED",
                error_message="Tavily API Key 无效或无权访问，请重新配置 Tavily Key。",
            )

    model = ScriptedModel([("search_sources", {
        "query": "测试公司 年报", "reason": "缺少年报",
    })])
    session = service.create(llm=model)
    session.data_source_preference = "online"
    service.store.save_research(session)
    service.attach_search(session.session_id, FailedSearch())
    state = service.turn(session.session_id, ResearchTurn(content="联网找年报"))
    assert len(model.calls) == 2
    assert state["session"]["search_history"][0]["status"] == "failed"
    assert state["session"]["facts"] == []


def test_ambiguous_multi_value_candidate_is_isolated_without_losing_valid_fact(service):
    runtime = session_runtime(service)
    args = fixture_observations(runtime, [{"metric": "revenue", "raw_value": "12000"},
                                          {"metric": "capital_expenditure", "raw_value": "100 90"}])
    state, model = run_turn(runtime, extraction_steps(args))
    assert [fact["standard_metric"] for fact in state["session"]["facts"]] == ["revenue"]
    assert "AMOUNT_NOT_EXPLICIT" in json.dumps(model.calls[1], ensure_ascii=False)
    assert state["session"]["facts"][0]["status"] == "confirmed"

def test_verified_facts_and_sources_survive_restart(service):
    runtime = session_runtime(service)
    args = fixture_observations(runtime)
    state, model = run_turn(runtime, [("read_file", {"file_id": args["file_id"], "view": "text"}), *extraction_steps(args)])
    fact = state["session"]["facts"][0]
    assert fact["status"] == "confirmed" and fact["normalized_value"] == "120000000"
    assert len(state["session"]["documents"][0]["sha256"]) == 64
    assert "PRIVATE_TEST_REASONING" not in json.dumps(state)
    assert any(message.get("reasoning_content") == "PRIVATE_TEST_REASONING" for message in model.calls[-1])
    fresh = ResearchService(SQLiteRunStore(service.store.data_dir))
    saved = fresh.store.get_research(runtime.session.session_id).facts[0]
    assert saved.status == "confirmed"
    assert saved.verification["semantic_review"]["status"] == "supported"
    assert fact["quote"] in fresh.store.research_blocks(runtime.session.session_id, args["file_id"])[0]["text"]
    assert not fresh.store.list_runs()

def test_long_conversation_keeps_explicit_durable_memory_after_restart(service):
    first = ScriptedModel([
        ("update_memory", {
            "updates": [{
                "key": "scope.statement_basis",
                "kind": "constraint",
                "content": "本任务后续分析统一采用合并报表口径；发现母公司口径时先询问用户。",
            }],
        }),
        ("finish_response", {
            "answer": "已记住本任务统一采用合并口径。",
        }),
    ])
    session = service.create(llm=first)
    service.turn(session.session_id, ResearchTurn(content="后续统一采用合并口径，遇到母公司口径先问我。"))
    for index in range(30):
        service.store.add_message(session.session_id, "user", f"临时讨论 {index}", "research")
        service.store.add_message(session.session_id, "assistant", f"临时回复 {index}", "research")

    fresh = ResearchService(SQLiteRunStore(service.store.data_dir))
    second = ScriptedModel([("finish_response", {"answer": "我会继续遵循已保存口径。"})])
    fresh.attach(session.session_id, second)
    state = fresh.turn(session.session_id, ResearchTurn(content="继续之前的任务"))

    system_prompt = second.calls[0][1]["content"]
    assert "scope.statement_basis" in system_prompt
    assert "统一采用合并报表口径" in system_prompt
    assert state["session"]["memory"][0]["source_message_id"].startswith("msg_")


def test_credentials_are_redacted_from_messages_memory_and_tool_events(service):
    secret = "sk-testsecret1234567890"
    model = ScriptedModel([
        ("update_memory", {"updates": [{"key": "preference.secret", "kind": "preference", "content": secret}]}),
        ("finish_response", {"answer": f"不会保存 {secret}"}),
    ])
    session = service.create(llm=model)
    state = service.turn(session.session_id, ResearchTurn(content=f"请记住 {secret}"))
    serialized = json.dumps(state, ensure_ascii=False)
    assert secret not in serialized
    assert "REDACTED_CREDENTIAL" in serialized
    assert state["session"]["memory"] == []
    assert any(event["type"] == "memory.updated" and event["payload"]["rejected_keys"] for event in state["events"])


def test_model_failure_can_recover_in_same_conversation(service):
    class FailingModel:
        def chat(self, messages, **kwargs):
            raise LlmError("LLM_HTTP_503: 供应商暂时不可用")
    session = service.create(llm=FailingModel())
    state = service.turn(session.session_id, ResearchTurn(content="继续整理资料"))
    assert state["session"]["last_issue"]["code"] == "LLM_HTTP_503"
    assert state["execution"]["status"] == "failed"
    service.attach(session.session_id, ScriptedModel([("finish_response", {"answer": "重试成功。"})]))
    state = service.turn(session.session_id, ResearchTurn(content="继续"))
    assert state["session"]["last_issue"] is None
    assert state["messages"][-1]["content"] == "重试成功。"


def test_registered_tool_provider_joins_same_audit_loop(tmp_path):
    class FinanceProbeProvider:
        provider_id = "finance-probe"
        version = "test-1"

        def tool_specs(self, session):
            return [ToolSpec(
                "inspect_finance_contract",
                "读取未来金融插件的能力契约，不执行估值。",
                NoArguments,
                lambda _: {"available": False, "reason": "正式模型待接入"},
            )]

    model = ScriptedModel([
        ("inspect_finance_contract", {}),
        ("finish_response", {"answer": "已检查金融插件契约；正式模型仍待接入。"}),
    ])
    service = ResearchService(
        SQLiteRunStore(tmp_path / "provider-runtime"),
        tool_providers=[FinanceProbeProvider()],
    )
    session = service.create(llm=model)
    state = service.turn(session.session_id, ResearchTurn(content="检查金融工具是否可用"))
    completed = [event for event in state["events"] if event["type"] == "tool.completed"]
    assert any(event["tool"] == "inspect_finance_contract" for event in completed)
    agent_start = next(event for event in state["events"] if event["type"] == "agent.started")
    assert agent_start["payload"]["tool_extensions"][0]["provider_id"] == "finance-probe"


def test_invented_number_is_rejected_and_tool_feedback_reaches_model(service):
    runtime = session_runtime(service)
    args = fixture_observations(runtime)
    args["rows"][0]["raw_value"] = "999999"
    result, model = run_turn(runtime, [("extract_observations", args),
        ("finish_response", {"answer": "原文不支持这个数值，需要重新核对。"})])
    assert not result["session"]["facts"]
    assert "AMOUNT_MISMATCH" in json.dumps(model.calls[-1], ensure_ascii=False)
    assert "原文不支持" in result["messages"][-1]["content"]
    assert any(event["tool"] == "extract_observations" and event["payload"].get("output", {}).get("ok") is False
               for event in result["events"])

@pytest.mark.parametrize("sign", ["-", "−"])
def test_extraction_cannot_drop_a_negative_source_sign(service, sign):
    from copy import deepcopy
    runtime = session_runtime(service)
    correct = fixture_observations(runtime, [{"metric": "net_income_parent", "raw_value": sign + "12000", "unit": "万元"}])
    wrong = deepcopy(correct)
    wrong["anchors"]["row0_value"]["quote"] = "12000"
    wrong["rows"][0]["raw_value"] = "12000"
    state, model = run_turn(runtime, [("extract_observations", wrong), *extraction_steps(correct)])
    assert len(state["session"]["facts"]) == 1
    assert state["session"]["facts"][0]["normalized_value"] == "-120000000"
    assert state["session"]["facts"][0]["status"] == "confirmed"
    assert "AMOUNT_BOUNDARY" in json.dumps(model.calls[1], ensure_ascii=False)

def test_extraction_understands_parenthetical_accounting_negatives(service):
    runtime = session_runtime(service)
    args = fixture_observations(runtime, [{"metric": "operating_receivables_decrease", "raw_value": "(1,200)", "unit": "万元", "period": "2024"}])
    state, _ = run_turn(runtime, extraction_steps(args))
    assert state["session"]["facts"][0]["raw_value"] == "(1,200)"
    assert state["session"]["facts"][0]["normalized_value"] == "-12000000"
    assert state["session"]["facts"][0]["status"] == "confirmed"

def test_repeated_identical_tool_failure_stops_without_fabricating(service):
    model = ScriptedModel([("missing_tool", {})] * 4)
    session = service.create(llm=model)
    state = service.turn(session.session_id, ResearchTurn(content="检查资料"))
    assert state["session"]["facts"] == []
    assert state["execution"]["status"] == "failed"
    assert state["session"]["last_issue"]["code"] == "AGENT_NO_PROGRESS"
    assert len(model.calls) == 2
    assert not service.store.list_runs()


def test_uncertain_fields_are_not_blanket_confirmed(service):
    runtime = session_runtime(service)
    args = fixture_observations(runtime)
    args["rows"][0]["uncertainties"] = ["原文单位适用范围尚未确定，不能凭模型自信确认"]
    result, model = run_turn(runtime, [*extraction_steps(args),
        ("finish_response", {"answer": "单位需要原文核验。", "outcome": "needs_input"})])
    assert len(result["session"]["facts"]) == 1
    assert result["session"]["facts"][0]["status"] == "proposed"
    assert result["session"]["facts"][0]["warnings"]
    assert "REVIEW_UNRESOLVED" in json.dumps(model.calls[-1], ensure_ascii=False)
    assert not service.store.list_runs()

def test_search_gap_does_not_fabricate_peers(service):
    model = ScriptedModel([
        ("search_sources", {"query": "可比公司", "reason": "可比公司不足"}),
        ("finish_response", {"answer": "搜索未配置，不能虚构同业数据。", "outcome": "insufficient_data"}),
    ])
    session = service.create(llm=model, data_source_preference="web")
    result = service.turn(session.session_id, ResearchTurn(content="补充同业"))
    assert result["session"]["facts"] == []
    assert result["session"]["search_history"][0]["status"] == "not_configured"
    assert result["session"]["outcome_status"] == "insufficient_data"


def test_search_provider_can_be_attached_to_one_session_without_persisting_key(service):
    from valuationagent.search.providers import MockSearchProvider

    session = service.create()
    provider = MockSearchProvider()
    service.attach_search(session.session_id, provider)
    assert service._search_clients[session.session_id] is provider
    event = service.store.list_events(session.session_id)[-1]
    assert event.type == "search.attached"
    assert event.payload == {"provider": "mock-search", "provider_version": "0.1"}


def test_search_provider_can_hot_attach_during_active_research_without_session_race(service):
    from valuationagent.search.providers import MockSearchProvider

    session = service.create()
    revision = session.revision
    assert service.store.acquire(session.session_id, "active-turn")
    try:
        provider = MockSearchProvider()
        service.attach_search(session.session_id, provider)
        status = service.data_service_status(session.session_id)
    finally:
        service.store.release(session.session_id, "active-turn")

    assert status["search"] == {
        "available": True,
        "provider": "mock-search",
        "connection_status": "configured",
    }
    assert service._search_client_epochs[session.session_id] == 1
    # Hot attachment is process-local and must not overwrite the active
    # turn's optimistic ResearchSession snapshot.
    assert service.store.get_research(session.session_id).revision == revision


def test_official_report_disclosure_total_is_deterministically_staged_not_dividend_base(service):
    session = service.create()
    session.draft.company = "美的集团"
    session.draft.ticker = "000333.SZ"
    session.draft.valuation_date = date(2025, 3, 31)
    session.information_cutoff_date = date(2025, 3, 31)
    file_id = "midea_official_report"
    session.documents.append(DocumentSummary(
        file_id=file_id,
        name="美的集团2024年年度报告.pdf",
        role="historical_financials",
        block_count=1,
        provenance_type="official_filing",
        authority_tier="A",
        provider="cninfo-announcements",
        source_url=(
            "https://static.cninfo.com.cn/finalpage/2025-03-29/1222951181.PDF"
        ),
    ))
    service.store.save_research_blocks(session.session_id, file_id, [{
        "file_id": file_id,
        "block_id": file_id + ":4",
        "text": (
            "美的集团股份有限公司2024年年度报告。"
            "以截至本报告披露之日公司总股本 7,660,355,772股"
            "扣除已回购股份28,452,226股后的股本总额7,631,903,546股为基数。"
        ),
        "location": {
            "page": 4,
            "source_type": "remote_document",
            "source_url": (
                "https://static.cninfo.com.cn/finalpage/2025-03-29/1222951181.PDF"
            ),
            "published_at": "2025-03-29",
        },
    }])

    candidates = service._deterministic_report_disclosure_shares(
        session, service._blocks(session), []
    )

    assert len(candidates) == 1
    assert candidates[0].normalized_value == "7660355772"
    assert candidates[0].period == "2025-03-29"
    assert candidates[0].scope == "issuer"
    assert candidates[0].verification["binding"] == "issuer_report_disclosure_shares"
    assert "7631903546" not in candidates[0].normalized_value


def test_clarification_is_free_conversation_without_numeric_approval(service):
    answer = "表头报告期无法可靠对应，请补充该页上下文。"
    model = ScriptedModel([("finish_response", {"answer": answer, "outcome": "needs_input"})])
    session = service.create(llm=model)
    result = service.turn(session.session_id, ResearchTurn(content="提取这份表格"))
    assert result["messages"][-1]["content"] == answer
    assert "question" not in result["session"]
    assert not result["session"]["facts"]
    assert not service.store.list_runs()


def test_large_fact_history_is_bounded_and_old_user_notes_are_retrievable(service):
    model = ScriptedModel([
        ("inspect_context", {"section": "facts", "query": "historical_revenue_000", "limit": 1}),
        ("inspect_context", {"section": "user_notes", "query": "早期成本口径"}),
        ("finish_response", {"answer": "已检索到早期字段和用户说明。"}),
    ])
    session = service.create(llm=model)
    session.facts = [FactCandidate(
        fact_id=f"fact_{index}", metric=f"historical_revenue_{index:03}",
        raw_value="100", unit="万元", normalized_value="1000000", period="2025",
        scope="consolidated", block_id=f"source:{index}", quote="营业收入 100 万元",
        status="confirmed",
    ) for index in range(620)]
    service.store.save_research(session)
    service.store.add_message(session.session_id, "user", "早期成本口径：统一使用主营业务成本。", "research")
    for index in range(30):
        service.store.add_message(session.session_id, "user", f"近期说明 {index}", "research")
    state = service.turn(session.session_id, ResearchTurn(content="请找回早期的字段和成本说明"))
    assert state["session"]["last_issue"] is None
    assert len(model.calls[0][1]["content"]) < 50000
    assert '"facts_omitted":' in model.calls[0][1]["content"]
    fact_output = json.loads(next(message["content"] for message in model.calls[1] if message["role"] == "tool"))
    assert fact_output["total"] == 1
    assert fact_output["items"][0]["fact_id"] == "fact_0"
    note_output = json.loads([message["content"] for message in model.calls[2] if message["role"] == "tool"][-1])
    assert note_output["items"][0]["text"] == "早期成本口径：统一使用主营业务成本。"
    assert note_output["items"][0]["block_id"].startswith("message:")


def test_context_retrieval_paginates_without_losing_history(service):
    model = ScriptedModel([
        ("inspect_context", {"section": "user_notes", "limit": 2}),
        ("inspect_context", {"section": "user_notes", "offset": 2, "limit": 2}),
        ("finish_response", {"answer": "已读取全部说明。"}),
    ])
    session = service.create(llm=model)
    for index in range(3):
        service.store.add_message(session.session_id, "user", f"历史说明 {index}", "research")
    service.turn(session.session_id, ResearchTurn(content="查看历史说明"))
    first = json.loads(next(message["content"] for message in model.calls[1] if message["role"] == "tool"))
    second = json.loads([message["content"] for message in model.calls[2] if message["role"] == "tool"][-1])
    assert first["next_offset"] == 2
    assert second["next_offset"] is None
    assert [item["text"] for item in first["items"] + second["items"]] == [
        "历史说明 0", "历史说明 1", "历史说明 2", "查看历史说明",
    ]


@pytest.mark.parametrize("response", [None, {"tool_calls": "invalid"}, {"tool_calls": [None]},
    {"tool_calls": [{"id": "call_bad", "function": None}]},
    {"tool_calls": [{"id": "call_bad", "function": {"name": [], "arguments": "{}"}}]}])
def test_malformed_model_protocol_becomes_actionable_recovery(service, response):
    class MalformedModel:
        def chat(self, messages, **kwargs):
            return response

    session = service.create(llm=MalformedModel())
    state = service.turn(session.session_id, ResearchTurn(content="继续整理资料"))
    assert state["session"]["last_issue"]["code"] == "TOOL_RESPONSE_INVALID"
    assert "question" not in state["session"]
    assert not any(event["type"] == "tool.started" for event in state["events"])


def test_model_errors_do_not_persist_credentials_in_recovery_or_audit(service):
    secret = "sk-testcredential123456789"

    class FailingModel:
        def chat(self, messages, **kwargs):
            raise LlmError("LLM_HTTP_400: 网关返回错误，凭证 " + secret)

    session = service.create(llm=FailingModel())
    state = service.turn(session.session_id, ResearchTurn(content="继续"))
    assert secret not in json.dumps(state, ensure_ascii=False)
    assert "REDACTED_CREDENTIAL" in state["session"]["last_issue"]["message"]


def test_tool_errors_and_question_labels_do_not_persist_credentials(service):
    secret = "sk-testcredential123456789"

    class FailingProvider:
        provider_id = "failing-provider"
        version = "test"

        def tool_specs(self, session):
            def fail(_):
                raise ValueError("上游拒绝凭证 " + secret)
            return [ToolSpec("failing_tool", "test", NoArguments, fail)]

    service.tool_providers = (FailingProvider(),)
    model = ScriptedModel([
        ("failing_tool", {}),
        ("finish_response", {"answer": "需要重新连接。不要保存 " + secret, "outcome": "needs_input"}),
    ])
    session = service.create(llm=model)
    state = service.turn(session.session_id, ResearchTurn(content="继续"))
    assert secret not in json.dumps(state, ensure_ascii=False)
    assert state["session"]["last_issue"] is None
    assert secret not in json.dumps(model.calls[-1], ensure_ascii=False)


def test_excel_locations_and_formulas_are_not_executed(service):
    book = Workbook()
    book.active.title = "历史财务"
    book.active.append(["营业收入", 12000, "万元"])
    book.active.append(["外部公式", '=WEBSERVICE("https://invalid.example")'])
    data = BytesIO()
    book.save(data)
    meta = service.store.save_upload("财务.xlsx", "historical_financials", None, data.getvalue())
    blocks, _ = parse_document(service.store.get_file(meta["file_id"]))
    assert blocks[0]["location"] == {"sheet": "历史财务", "row": 1}
    assert "B1: 12000" in blocks[0]["text"]
    assert "公式，未求值" in blocks[1]["text"]
    assert "https://" not in blocks[1]["text"]


def test_docx_upload_preserves_paragraph_and_table_order(service):
    from docx import Document

    document = Document()
    document.add_heading("资本开支假设", level=1)
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "参数"
    table.cell(0, 1).text = "数值"
    table.cell(1, 0).text = "alpha"
    table.cell(1, 1).text = "1.0"
    document.add_paragraph("表后说明")
    data = BytesIO()
    document.save(data)

    meta = service.store.save_upload(
        "估值方案.docx", "evidence", None, data.getvalue()
    )
    blocks, warnings = parse_document(service.store.get_file(meta["file_id"]))

    assert warnings == []
    assert [block["text"] for block in blocks] == [
        "资本开支假设",
        "参数 | 数值",
        "alpha | 1.0",
        "表后说明",
    ]
    assert blocks[0]["location"]["paragraph"] == 1
    assert blocks[1]["location"] == {"table": 1, "row": 1, "cells": ["参数", "数值"]}
    assert blocks[3]["location"]["paragraph"] == 2


def test_api_workspace_sources_are_scoped_and_exports_are_safe(tmp_path):
    with TestClient(create_app(tmp_path)) as client:
        first = client.post("/api/workspaces", json={"language": "en-US"}).json()["workspace"]["workspace_id"]
        second = client.post("/api/workspaces", json={}).json()["workspace"]["workspace_id"]
        fid = client.post("/api/files", data={"role": "evidence"}, files={"file": ("policy.txt", b"<script>alert(1)</script>", "text/plain")}).json()["file_id"]
        assert client.post(f"/api/workspaces/{first}/messages", json={"file_ids": [fid]}).status_code == 202
        result = client.get(f"/api/workspaces/{first}").json()
        assert result["research"]["session"]["documents"][0]["name"] == "policy.txt"
        assert client.get(f"/api/workspaces/{first}/sources/{fid}").status_code == 200
        assert client.get(f"/api/workspaces/{second}/sources/{fid}").status_code == 404
        assert client.post(f"/api/workspaces/{first}/messages", json={}).status_code == 422
        report = client.get(f"/api/workspaces/{first}/export?format=json").json()
        assert not report["result_document"]["numeric_result_available"]
        html = client.get(f"/api/workspaces/{first}/export?format=html").text
        assert "<script>" not in html and "No formal valuation submitted" in html


@pytest.mark.parametrize("thinking,expected", [("auto", "required"), ("enabled", "auto")])
def test_deepseek_tool_payload_and_redacted_diagnostics(thinking, expected):
    recorded = []
    secret = "TEST_KEY_NOT_REAL"

    def handler(request):
        recorded.append(json.loads(request.content))
        return httpx.Response(400, json={"error": {"message": "tool_choice rejected " + secret}})

    original = httpx.Client
    config = ModelConnectionInput(base_url="https://api.deepseek.com", model="deepseek-v4-flash", api_key=secret, thinking=thinking)
    with (
        patch(
            "valuationagent.llm.client.httpx.Client",
            side_effect=lambda **kwargs: original(
                transport=httpx.MockTransport(handler), **kwargs
            ),
        ),
        pytest.raises(Exception) as caught,
    ):
        OpenAICompatibleClient(config).chat(
            [{"role": "user", "content": "hello"}],
            tools=[{"type": "function"}],
            tool_choice="required",
        )
    assert recorded[0]["tool_choice"] == expected
    assert "temperature" not in recorded[0]
    assert "tool_choice" in str(caught.value) and secret not in str(caught.value)


def test_llm_timeout_and_invalid_json_have_distinct_recovery_codes():
    original = httpx.Client
    config = ModelConnectionInput(
        base_url="https://example.test/v1", model="test-model", api_key="TEST_ONLY"
    )

    def timeout_handler(request):
        raise httpx.ReadTimeout("late", request=request)

    with (
        patch(
            "valuationagent.llm.client.httpx.Client",
            side_effect=lambda **kwargs: original(
                transport=httpx.MockTransport(timeout_handler), **kwargs
            ),
        ),
        pytest.raises(LlmError, match="LLM_TIMEOUT"),
    ):
        OpenAICompatibleClient(config).chat([{"role": "user", "content": "hello"}])

    def invalid_json_handler(request):
        return httpx.Response(200, content=b"not-json")

    with (
        patch(
            "valuationagent.llm.client.httpx.Client",
            side_effect=lambda **kwargs: original(
                transport=httpx.MockTransport(invalid_json_handler), **kwargs
            ),
        ),
        pytest.raises(LlmError, match="LLM_RESPONSE_INVALID_JSON"),
    ):
        OpenAICompatibleClient(config).chat([{"role": "user", "content": "hello"}])
