"""Synthetic, offline acceptance of continuous modeling, not live-LLM accuracy."""
from decimal import Decimal
from io import BytesIO

import pytest
from openpyxl import load_workbook
from pydantic import ValidationError
from test_evidence_recovery import block, configured_service, item
from test_finance_team_model import history
from test_research_sessions import ScriptedModel

from valuationagent.application.agent_runtime import WorkspaceAgentRuntime
from valuationagent.application.reporting import ValuationReportExporter
from valuationagent.application.reproducibility import (
    build_valuation_bundle,
    replay_bundle,
)
from valuationagent.application.research import ProposeFacts, ProposeForecast
from valuationagent.application.runner import ValuationRunner
from valuationagent.application.valuation_plan import valuation_progress
from valuationagent.finance.team_model import FinanceTeamModel
from valuationagent.schemas.models import required_financial_metrics
from valuationagent.schemas.research import (
    DocumentSummary,
    FactCandidate,
    ForecastInputs,
    ResearchTurn,
)


def forecast(evidence_id="file_test:1"):
    return {"inputs": {"revenue_growth_scenarios": {
        "pessimistic": ["0.01"] * 10, "base": ["0.05"] * 10, "optimistic": ["0.08"] * 10},
        "ebit_margin_scenarios": {"pessimistic": ["0.18"] * 10, "base": ["0.24"] * 10, "optimistic": ["0.27"] * 10},
        "wacc": "0.10", "terminal_growth": "0.02"},
        "rationale": "这是合成测试的预测判断，不是历史事实；以已核验基期为锚，设置不同需求增长及利润率路径，并明确折现率与终值假设。",
        "risks": ["历史样本仅一年，无法验证增长趋势；预测及折现率存在主观性。"],
        "evidence_ids": [evidence_id]}


def baseline(service, session, *, omit=(), methods=("dcf",)):
    row = history()[-1]
    fields = sorted(required_financial_metrics(methods) - set(omit))
    units = {key: "ratio" if key in {"ebit_margin", "tax_rate"} else "股" if key == "common_shares" else "元" for key in fields}
    lines = {key: f"{key}（{units[key]}） {getattr(row, key)}" for key in fields}
    source = block(f"{session.draft.company} {session.draft.ticker} 2025年合并报表\n单位：元\n" + "\n".join(lines.values()))
    meta = service.store.save_upload("synthetic-baseline.txt", "historical_financials", "text/plain", source["text"].encode())
    fid = meta["file_id"]
    source.update({"block_id": fid + ":1", "file_id": fid})
    session.documents = [DocumentSummary(file_id=fid, name=meta["original_name"], role="historical_financials",
                                        block_count=1, sha256=meta["sha256"], size_bytes=meta["size_bytes"])]
    service.store.save_research_blocks(session.session_id, fid, [source])
    return [item(key, str(getattr(row, key)), block_id=fid + ":1", unit=units[key], quote=lines[key]) for key in fields]


@pytest.mark.parametrize("company,ticker", [("样本制造甲公司", "600123"), ("样本制造乙公司", "000321"),
                                           ("样本制造丙公司", "300123"), ("样本制造丁公司", "688123")])
def test_multiple_batches_to_one_review_to_dcf_sensitivity_reports_and_replay(tmp_path, company, ticker):
    from valuationagent.application.workspaces import ValuationWorkspaceService
    from valuationagent.schemas.workspace import ValuationWorkspace
    service, session, _ = configured_service(tmp_path)
    session.draft.company, session.draft.ticker = company, ticker
    session.documents = []
    from observation_fixtures import fixture_observations, ObservationModel, extraction_steps
    runtime = WorkspaceAgentRuntime(service, session)
    values = history()[-1]
    specifications = []
    for metric in sorted(required_financial_metrics(["dcf"])):
        unit = "ratio" if metric in {"ebit_margin", "tax_rate"} else "股" if metric == "common_shares" else "元"
        specification = {"metric": metric, "raw_value": str(getattr(values, metric)), "unit": unit}
        if metric == "common_shares":
            specification["period"] = "2026-06-30"
        if metric in {"interest_bearing_debt", "cash_and_non_operating_assets"}:
            specification.update(semantic_role="financing" if metric == "interest_bearing_debt" else "non_operating",
                                 ebit_treatment="exclude", fcff_treatment="exclude", equity_bridge_treatment="include")
        specifications.append(specification)
    first = fixture_observations(runtime, specifications[:5])
    second = fixture_observations(runtime, specifications[5:])
    model = ObservationModel([
        ("update_task", {"draft": session.draft.model_dump(mode="json"), "valuation_requested": True}),
        *extraction_steps(first),
        *extraction_steps(second),
        ("propose_forecast", forecast(first["file_id"] + ":1")),
        ("calculate_valuation", {}),
        ("finish_response", {"answer": "已完成确定性计算，预测是明确标注的假设。"}),
    ])
    service._clients[session.session_id] = model
    service.store.save_research(session)
    workspaces = ValuationWorkspaceService(service.store, service, ValuationRunner(service.store, FinanceTeamModel()))
    workspace = ValuationWorkspace(workspace_id="workspace_" + ticker, research_session_id=session.session_id)
    service.store.create_workspace(workspace)
    workspaces.message(workspace.workspace_id, ResearchTurn(content="开始自动化估值"))
    state = workspaces.snapshot(workspace.workspace_id)
    assert len(model.calls) == 10
    assert "question" not in state["research"]["session"]
    assert {fact["status"] for fact in state["research"]["session"]["facts"]} == {"confirmed"}
    assert state["research"]["session"]["forecast_proposal"]["status"] == "confirmed"
    result = service.store.get_run(state["workspace"]["active_run_id"])
    assert str(result.status).startswith("completed"), result.model_dump(mode="json")
    assert result.result.dcf and result.result.sensitivity and result.result.sensitivity_studies
    assert len(result.result.forecast) == 10
    assert result.request.financials.revenue == history()[-1].revenue
    assert "模型推断" in result.request.assumption_evidence["wacc"][0].note
    assert replay_bundle(build_valuation_bundle(service.store, result))["passed"]
    report = ValuationReportExporter()
    assert "预测与FCFF" in load_workbook(BytesIO(report.xlsx(result))).sheetnames
    assert report.pdf(result).startswith(b"%PDF")


def workspace_for(service, session):
    from valuationagent.application.workspaces import ValuationWorkspaceService
    from valuationagent.schemas.workspace import ValuationWorkspace
    service.store.save_research(session)
    workspaces = ValuationWorkspaceService(service.store, service, ValuationRunner(service.store, FinanceTeamModel()))
    workspace = ValuationWorkspace(workspace_id="workspace_" + session.session_id, research_session_id=session.session_id, run_policy="review")
    service.store.create_workspace(workspace)
    return workspaces, workspace


def staged_model(tmp_path, omit=()):
    service, session, _ = configured_service(tmp_path)
    session.pending_action = "valuation"
    result = WorkspaceAgentRuntime(service, session).facts(ProposeFacts(candidates=baseline(service, session, omit=omit)))
    assert result["candidates"] and not result.get("_terminal")
    service._propose_forecast(session, ProposeForecast.model_validate(forecast(session.documents[0].file_id + ":1")))
    return service, session


def test_forecast_never_fills_missing_historical_debt_or_approves_facts(tmp_path):
    service, session = staged_model(tmp_path, omit=("interest_bearing_debt",))
    progress = valuation_progress(session, service.valuation_assembler)
    assert not progress["ready_for_review"]
    assert "有息" in progress["blocking_reason"]
    assert all(f.status == "confirmed" for f in session.facts)
    assert session.forecast_proposal.status == "proposed"
    assert not progress["ready_for_review"]
    assert "question" not in session.model_dump()


def test_bound_capex_candidate_requires_semantic_repair_before_fallback_review(tmp_path):
    service, session = staged_model(tmp_path, omit=("capital_expenditure",))
    candidate = item(
        "购建固定资产、无形资产和其他长期资产支付的现金",
        "150",
        standard_metric="cash_paid_for_ppe_intangibles",
        semantic_role="unknown",
        ebit_treatment="exclude",
        fcff_treatment="include",
        equity_bridge_treatment="exclude",
        mapping_confidence=0.9,
        mapping_rationale="合并现金流量表投资活动中的长期资产购建现金支出，数值及年度列已完成绑定。",
    )
    warned = FactCandidate(
        **candidate.model_dump(),
        fact_id="repairable_capex",
        normalized_value="150",
        warnings=[
            "语义映射缺少经营/融资/金融子公司等经济角色",
            "资本开支基础科目模型处理冲突：投资活动现金支出应标记investing，作为FCFF确定性推导输入，不进入EBIT或权益桥接",
        ],
    )
    session.facts.append(warned)

    progress = valuation_progress(session, service.valuation_assembler)

    assert progress["status"] == "quality_repair_required"
    assert not progress["ready_for_review"]
    assert progress["recoverable_driver_repairs"][0]["required_mapping"]["semantic_role"] == "investing"
    assert "不能在可修复时静默改用比例回退" in progress["blocking_reason"]


def test_single_pe_method_still_requires_peer_multiples(tmp_path):
    service, session, _ = configured_service(tmp_path)
    session.draft.methods = ["pe"]
    session.pending_action = "valuation"
    result = WorkspaceAgentRuntime(service, session).facts(ProposeFacts(candidates=baseline(service, session, methods=("pe",))),
    )
    assert result["candidates"]
    progress = valuation_progress(session, service.valuation_assembler)
    assert not progress["ready_for_review"]
    assert "PE 当前0家" in progress["blocking_reason"]
    assert "目标公司自身收盘价" in progress["blocking_reason"]
    assert "question" not in session.model_dump()


def test_missing_peer_data_does_not_block_a_ready_dcf(tmp_path):
    service, session, _ = configured_service(tmp_path)
    session.draft.methods = ["dcf", "pe"]
    session.pending_action = "valuation"
    WorkspaceAgentRuntime(service, session).facts(ProposeFacts(candidates=baseline(service, session, methods=("dcf", "pe"))))
    service._propose_forecast(session, ProposeForecast.model_validate(forecast(session.documents[0].file_id + ":1")))
    workspaces, workspace = workspace_for(service, session)
    review = workspaces.prevaluation_review(workspace.workspace_id)
    assert review.executable_methods == ["dcf"] and "pe" in review.excluded_methods
    record = workspaces.approve(workspace.workspace_id, review.checkpoint_id)
    result = workspaces.execute(record.run_id)
    assert result.result and result.result.dcf
    assert result.request.methods == ["dcf"] and "pe" in result.request.excluded_methods
    assert any("PE因可靠数据不足未进入本次计算" in warning for warning in result.result.warnings)


def test_partial_older_history_does_not_block_reviewed_explicit_forecast(tmp_path):
    service, session = staged_model(tmp_path)
    old = session.facts[0].model_copy(deep=True, update={"fact_id": "old_partial", "period": "2023"})
    session.facts.append(old)
    progress = valuation_progress(session, service.valuation_assembler)
    assert progress["ready_for_review"] and progress["historical_periods"] == []
    assert old.status == "confirmed"


@pytest.mark.parametrize("change", ["short_path", "missing_scenario", "percent_not_ratio", "nan", "bad_order", "bad_terminal"])
def test_invalid_forecasts_are_rejected_before_staging(change):
    values = forecast()["inputs"]
    if change == "short_path": values["revenue_growth_scenarios"]["base"].pop()
    if change == "missing_scenario": del values["revenue_growth_scenarios"]["pessimistic"]
    if change == "percent_not_ratio": values["revenue_growth_scenarios"]["base"][0] = "8"
    if change == "nan": values["revenue_growth_scenarios"]["base"][0] = "NaN"
    if change == "bad_order": values["revenue_growth_scenarios"]["pessimistic"][0] = "0.20"
    if change == "bad_terminal": values["terminal_growth"] = "0.10"
    with pytest.raises(ValidationError): ForecastInputs.model_validate(values)


def test_stale_plan_cannot_confirm_values_or_forecast_after_scope_changes(tmp_path):
    service, session = staged_model(tmp_path)
    workspaces, workspace = workspace_for(service, session)
    review = workspaces.prevaluation_review(workspace.workspace_id)
    session = service.store.get_research(session.session_id)
    service._apply_task_draft(session, session.draft.model_copy(update={"methods": ["pe"]}), automatic=True)
    service.store.save_research(session)
    with pytest.raises(ValueError, match="变化|阻塞"):
        workspaces.approve(workspace.workspace_id, review.checkpoint_id)
    assert not service.store.list_runs()


def test_reopening_checkpoint_does_not_calculate(tmp_path):
    service, session = staged_model(tmp_path)
    workspaces, workspace = workspace_for(service, session)
    review = workspaces.prevaluation_review(workspace.workspace_id)
    workspaces.reopen(workspace.workspace_id, review.checkpoint_id, "修改方案")
    assert workspaces.get(workspace.workspace_id).active_checkpoint_id is None
    assert service.store.get_research(session.session_id).forecast_proposal.status == "proposed"
    assert not service.store.list_runs()


def test_forecast_tool_rejects_unknown_and_search_snippet_evidence(tmp_path):
    service, session, _ = configured_service(tmp_path)
    session.pending_action = "valuation"
    source = block("只有搜索摘要", 2, source_type="web_search")
    service.store.save_research_blocks(session.session_id, "file_test", [source])
    for key in ("unknown", source["block_id"]):
        args = {**forecast(), "evidence_ids": [key]}
        with pytest.raises(ValueError, match="引用或搜索摘要"):
            service._propose_forecast(session, ProposeForecast.model_validate(args))
    assert session.forecast_proposal is None


def test_forecast_cannot_start_valuation_for_a_research_only_request(tmp_path):
    service, session, _ = configured_service(tmp_path)
    with pytest.raises(ValueError, match="已请求DCF"):
        service._propose_forecast(session, ProposeForecast.model_validate(forecast()))
    assert not session.pending_action and session.forecast_proposal is None


def test_confirmed_method_change_invalidates_old_forecast_without_blocking_new_method(tmp_path):
    service, session = staged_model(tmp_path)
    old_id = session.forecast_proposal.proposal_id
    revised = session.draft.model_copy(update={"methods": ["pe"]})
    service._apply_task_draft(session, revised, automatic=True)
    assert session.draft.methods == ["pe"] and session.forecast_proposal is None
    event = next(event for event in service.store.list_events(session.session_id) if event.type == "valuation.forecast_invalidated")
    assert event.payload["proposal_id"] == old_id


def test_verified_correction_changes_working_inputs_but_does_not_calculate(tmp_path):
    service, session, _ = configured_service(tmp_path)
    WorkspaceAgentRuntime(service, session).facts(ProposeFacts(candidates=baseline(service, session)))
    old_fact = next(fact for fact in session.facts if fact.metric == "common_shares")
    file_id = session.documents[0].file_id
    sources = service.store.research_blocks(session.session_id, file_id)
    replacement = {"block_id": file_id + ":2", "text": f"{session.draft.company} {session.draft.ticker} 2025年合并报表\ncommon_shares（股） 421000000", "location": {}}
    service.store.save_research_blocks(session.session_id, file_id, [*sources, replacement])
    candidate = item("common_shares", "421000000", unit="股", block_id=file_id + ":2", quote="common_shares（股） 421000000")
    result = WorkspaceAgentRuntime(service, session).facts(ProposeFacts(candidates=[candidate], replaces=[old_fact.fact_id]))
    assert result["candidates"]
    assert old_fact.status == "rejected" and session.facts[-1].status == "confirmed"
    assert session.staged_supersessions[session.facts[-1].fact_id] == [old_fact.fact_id]
    assert not service.store.list_runs()


def test_future_forecast_evidence_is_rejected(tmp_path):
    service, session, _ = configured_service(tmp_path)
    session.pending_action = "valuation"
    source = block("尚未披露的预测依据", published_at="2027-01-01")
    service.store.save_research_blocks(session.session_id, "file_test", [source])
    with pytest.raises(ValueError, match="未来信息"):
        service._propose_forecast(session, ProposeForecast.model_validate(forecast()))


def test_revising_forecast_over_confirmed_baseline_is_a_review_state_not_a_crash(tmp_path):
    service, session = staged_model(tmp_path)
    workspaces, workspace = workspace_for(service, session)
    original = workspaces.prevaluation_review(workspace.workspace_id)
    session = service.store.get_research(session.session_id)
    proposal = forecast(session.documents[0].file_id + ":1")
    proposal["inputs"]["wacc"] = "0.11"
    service._propose_forecast(session, ProposeForecast.model_validate(proposal))
    service.store.save_research(session)
    updated = workspaces.prevaluation_review(workspace.workspace_id)
    assert updated.state_hash != original.state_hash
    assert Decimal(updated.request_snapshot["assumptions"]["wacc"]) == Decimal("0.11")
    assert not service.store.list_runs()
